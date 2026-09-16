#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""build-trade-kb.py — 生成「外贸人知识库」内置资源包。

把源目录里的一批外贸业务表格 / 文档模板，转成 GeekClaw 前端可直接消费的静态资源：

    ui/public/trade-kb/
      manifest.json          分类 + 文件清单 + 元数据（前端据此渲染目录与卡片）
      files/<原始文件名>      原始文件，供用户「下载使用」（保持原文件名与格式）
      preview/<id>.html      自包含的「在线查阅」预览页（纯 HTML + 内联 CSS，无脚本）

预览页由构建期一次性生成，运行时不依赖任何外部工具（不调用 officecli / Office），
离线可用。Excel 走 xlrd(\.xls) / openpyxl(\.xlsx)，Word 走 python-docx(\.docx) 与
自实现的 Word97 分片表解析(\.doc)，PPT 走 python-pptx(\.pptx) 与自实现的
PowerPoint 二进制文本原子扫描(\.ppt)。

用法：
    <venv>/python.exe scripts/build-trade-kb.py \
        --source "E:/BaiduNetdiskDownload/外贸人的知识库/外贸人知识库" \
        --out    "ui/public/trade-kb"

依赖：xlrd openpyxl python-docx python-pptx olefile
"""

from __future__ import annotations

import argparse
import datetime as _dt
import hashlib
import html
import json
import os
import re
import shutil
import sys
import warnings

warnings.filterwarnings('ignore')

import xlrd  # noqa: E402
import openpyxl  # noqa: E402
import olefile  # noqa: E402
from docx import Document  # noqa: E402
from docx.table import Table as DocxTable  # noqa: E402
from docx.text.paragraph import Paragraph as DocxParagraph  # noqa: E402
from pptx import Presentation  # noqa: E402

# --------------------------------------------------------------------------- #
# 分类定义（顺序即左侧目录顺序）
# --------------------------------------------------------------------------- #

CATEGORIES: list[dict] = [
    {'id': 'customer', 'name': '客户档案管理', 'icon': 'People', 'order': 1},
    {'id': 'followup', 'name': '客户跟进体系', 'icon': 'Timeline', 'order': 2},
    {'id': 'customer-dev', 'name': '客户开发与话术', 'icon': 'Communication', 'order': 3},
    {'id': 'market', 'name': '市场调研与分析', 'icon': 'ChartLine', 'order': 4},
    {'id': 'exhibition', 'name': '展会与海外拜访', 'icon': 'Plane', 'order': 5},
    {'id': 'marketing', 'name': '平台营销推广', 'icon': 'Megaphone', 'order': 6},
    {'id': 'team', 'name': '团队与绩效管理', 'icon': 'PeopleFolder', 'order': 7},
    {'id': 'docs', 'name': '外贸单证模板', 'icon': 'FileText', 'order': 8},
    {'id': 'reference', 'name': '工具与使用说明', 'icon': 'Tool', 'order': 9},
]

# 文件名 → (分类, 一句话说明)。说明用于卡片与搜索，写清「这份模板解决什么问题」。
FILE_META: dict[str, tuple[str, str]] = {
    # —— 客户档案管理 ——
    '01 成交客户表.xls': ('customer', '成交客户总台账：记录客户信息、成交产品、订单金额与返单情况，作为客户资产的基础档案。'),
    '02 成交客户重要信息收集表.xls': ('customer', '成交客户深度背调表：收集决策人、采购习惯、付款方式、竞品供应商等关键信息。'),
    '03 送样客户表.xls': ('customer', '送样登记表 + 送样汇总表：跟踪样品寄送、反馈与转化结果，避免样品打水漂。'),
    '04 意向客户表.xls': ('customer', '意向客户跟进表：记录客户需求、意向等级与跟进动作，推动意向向订单转化。'),
    '05 询盘客户表.xls': ('customer', '询盘登记表：按来源、产品、需求量归集询盘，用于判断询盘质量与响应优先级。'),
    '06 流失客户表.xls': ('customer', '流失客户登记表：记录流失原因与最后接触时间，为后续二次激活提供依据。'),
    '07 意向订单统计表.xls': ('customer', '意向订单统计：汇总在谈订单的客户、金额与预计成交时间，用于预测业绩。'),
    '08 老客户返单分析表.xls': ('customer', '客户返单分析：统计老客户复购频次、金额与品类，识别高价值返单机会。'),
    '09 客户投诉登记表.xls': ('customer', '客户投诉登记表：记录投诉内容、责任归属、处理结果与闭环时间，沉淀质量改进依据。'),
    '21 样品记录模版.xlsx': ('customer', '样品记录模板：逐条登记样品编号、规格、寄送信息与客户反馈。'),
    # —— 客户跟进体系 ——
    '客户跟进表.xls': ('followup', '完整的客户跟进体系表格：含「客户总览」大表 + 500 张单客户跟进明细页，可按客户逐个建档跟踪。'),
    '国际站客户跟进表.xls': ('followup', '国际站（阿里国际站等）专用客户跟进体系：客户总览 + 500 张单客户页，适配平台询盘跟进节奏。'),
    '客户跟进表一览表.xlsx': ('followup', '客户跟进一览表：集中查看客户档案资料与跟进状态，作为跟进表体系的首页索引。'),
    # —— 客户开发与话术 ——
    '外贸开发客户各话术模板.xlsx': ('customer-dev', '外贸客户开发话术全集（18 个工作表）：覆盖领英档案与话术、LinkedIn/WhatsApp/Facebook 开发、Cold Call 电话 SOP、搜索指令、客户背调、价格谈判、展会跟进、自动回复邮件等。'),
    # —— 市场调研与分析 ——
    '10 市场客户需求调研表.xlsx': ('market', '市场客户需求调研表：结构化收集目标市场需求、采购偏好与价格带，输出选品与定位依据。'),
    '14 国外市场销售量及大客户统计.xls': ('market', '国外市场销售量与大客户统计：按市场汇总销量并识别头部客户，支撑区域策略。'),
    # —— 展会与海外拜访 ——
    '11 国外实地拜访客户计划.xls': ('exhibition', '国外实地拜访计划：规划拜访对象、目的、行程与预期成果。'),
    '12 展会计划.xls': ('exhibition', '展会参展计划表：展位、预算、物料、人员分工与目标客户一表规划。'),
    '13 国外拜访或参展费用预算.xls': ('exhibition', '海外拜访 / 参展费用预算表：差旅、展位、物料等成本项预算与实报对比。'),
    '23 国外拜访总结 模版.pptx': ('exhibition', '国外拜访总结汇报模板：行程复盘、客户反馈、市场发现与后续动作，可直接改为汇报 PPT。'),
    '24 国外拜访总结与市场调研报告 模板.docx': ('exhibition', '海外拜访总结与市场调研报告模板：把走访见闻整理成结构化市场情报报告。'),
    '25 展会总结报告 模版.ppt': ('exhibition', '展会总结报告模板：展中获客数据、客户质量、投入产出与改进建议。'),
    '26 展会客户表.xls': ('exhibition', '展会客户登记表：现场快速登记名片与需求，展后按意向分级跟进。'),
    # —— 平台营销推广 ——
    '15 付费平台推广计划.xls': ('marketing', '付费平台推广计划表：按平台规划投放预算、关键词与目标询盘量。'),
    '16 付费平台效果统计.xls': ('marketing', '付费平台效果统计表：对比各平台曝光、点击、询盘与成本，评估投放 ROI。'),
    '17 新客户来源统计表.xls': ('marketing', '新客户来源统计表：归因各渠道获客数量与质量，识别高效获客渠道。'),
    # —— 团队与绩效管理 ——
    '18 团队培训计划表.xls': ('team', '团队培训计划表：按岗位安排培训主题、时间与考核方式。'),
    '19 离职人员统计分析表.xls': ('team', '离职人员统计分析表：汇总离职原因、岗位与在职时长，辅助团队稳定性诊断。'),
    '20 年度重要工作计划推进表.xls': ('team', '年度重点工作推进表：把全年关键任务拆到月度并跟踪推进状态。'),
    '22 KPI绩效考核表.doc': ('team', '外贸业务 KPI 绩效考核表：量化业绩、过程与能力指标，可直接用于季度考核。'),
    # —— 外贸单证模板 ——
    '装箱单、商业发票、PI模板.xls': ('docs', '外贸三大基础单证模板：PI（形式发票）、Commercial Invoice（商业发票）、Packing List（装箱单），含多种版式，填好即可发客户。'),
    # —— 工具与使用说明 ——
    '上班时间表节假日表.xlsx': ('reference', '各国上班时间与节假日查询表：安排跨时区沟通、催单与交期时避免踩到客户假期。'),
    '表格重要项目说明.doc': ('reference', '全套表格使用说明：逐份说明各表用途、填写要点与彼此之间的数据衔接关系，建议先读这份。'),
}

# 预览生成上限（防止个别超宽表把 HTML 撑爆）
MAX_SHEETS = 12
MAX_ROWS = 600
MAX_COLS = 60

# --------------------------------------------------------------------------- #
# 通用小工具
# --------------------------------------------------------------------------- #

_CTRL_RE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def esc(v) -> str:
    return html.escape(str(v), quote=True)


def clean_text(s: str) -> str:
    return _CTRL_RE.sub('', s).strip()


def cell_to_text(v, datemode: int | None = None, ctype: int | None = None) -> str:
    """把单元格值格式化成适合展示的字符串。"""
    if v is None:
        return ''
    if ctype is not None and ctype == xlrd.XL_CELL_DATE and isinstance(v, float):
        try:
            dt = xlrd.xldate_as_datetime(v, datemode or 0)
            if dt.hour or dt.minute or dt.second:
                return dt.strftime('%Y-%m-%d %H:%M')
            return dt.strftime('%Y-%m-%d')
        except Exception:
            return str(v)
    if isinstance(v, _dt.datetime):
        if v.hour or v.minute or v.second:
            return v.strftime('%Y-%m-%d %H:%M')
        return v.strftime('%Y-%m-%d')
    if isinstance(v, _dt.date):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        return f'{v:g}'
    if isinstance(v, bool):
        return '是' if v else '否'
    return clean_text(str(v))


def is_blank_row(cells: list[str]) -> bool:
    return all(not c.strip() for c in cells)


def trim_grid(grid: list[list[str]]) -> list[list[str]]:
    """去掉尾部全空的行与列。"""
    while grid and is_blank_row(grid[-1]):
        grid.pop()
    if not grid:
        return []
    width = 0
    for row in grid:
        for i in range(len(row) - 1, -1, -1):
            if row[i].strip():
                width = max(width, i + 1)
                break
    return [row[:width] for row in grid]


# --------------------------------------------------------------------------- #
# Excel → HTML
# --------------------------------------------------------------------------- #

def sheet_to_html(grid: list[list[str]], merges: list[tuple[int, int, int, int]] | None,
                  truncated: bool) -> str:
    """grid 为已格式化的二维文本；merges 为 (r1, r2, c1, c2) 0-based 闭区间。"""
    if not grid:
        return "<p class='empty-sheet'>（此工作表为空）</p>"

    n_rows = len(grid)
    n_cols = max(len(r) for r in grid) if grid else 0

    skip: set[tuple[int, int]] = set()
    span: dict[tuple[int, int], tuple[int, int]] = {}
    for (r1, r2, c1, c2) in (merges or []):
        if r2 < r1 or c2 < c1:
            continue
        if r1 >= n_rows or c1 >= n_cols:
            continue
        r2 = min(r2, n_rows - 1)
        c2 = min(c2, n_cols - 1)
        span[(r1, c1)] = (r2 - r1 + 1, c2 - c1 + 1)
        if r2 == r1 and c2 == c1:
            continue
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                if (rr, cc) != (r1, c1):
                    skip.add((rr, cc))

    out: list[str] = ["<table class='sheet-table'>"]
    for r in range(n_rows):
        out.append('<tr>')
        for c in range(n_cols):
            if (r, c) in skip:
                continue
            text = grid[r][c] if c < len(grid[r]) else ''
            rs, cs = span.get((r, c), (1, 1))
            attrs = ''
            if rs > 1:
                attrs += f' rowspan="{rs}"'
            if cs > 1:
                attrs += f' colspan="{cs}"'
            cls = ' class="c-head"' if r == 0 else ''
            if text:
                out.append(f'<td{cls}{attrs}>{esc(text)}</td>')
            else:
                out.append(f'<td{cls}{attrs}></td>')
        out.append('</tr>')
    out.append('</table>')
    if truncated:
        out.append("<p class='trunc-note'>已截断展示，完整内容请下载原文件查看。</p>")
    return ''.join(out)


def read_xls_sheets(path: str) -> list[tuple[str, str]]:
    """返回 [(sheet 名, HTML)]。"""
    # formatting_info 才能拿到合并单元格；它与 on_demand 互斥，文件都很小，直接全量加载。
    has_fmt = True
    try:
        wb = xlrd.open_workbook(path, formatting_info=True)
    except Exception:
        wb = xlrd.open_workbook(path, on_demand=True)
        has_fmt = False

    names = wb.sheet_names()
    keep = min(len(names), MAX_SHEETS)
    result: list[tuple[str, str]] = []
    for idx in range(keep):
        name = names[idx]
        sh = wb.sheet_by_name(name)
        n_rows = min(sh.nrows, MAX_ROWS)
        n_cols = min(sh.ncols, MAX_COLS)
        grid: list[list[str]] = []
        for r in range(n_rows):
            row = [
                cell_to_text(sh.cell_value(r, c), wb.datemode, sh.cell_type(r, c))
                for c in range(n_cols)
            ]
            grid.append(row)
        grid = trim_grid(grid)
        merges = None
        if has_fmt:
            try:
                merges = list(sh.merged_cells)  # (rlo, rhi, clo, chi) 半开区间
                merges = [(r1, r2 - 1, c1, c2 - 1) for (r1, r2, c1, c2) in merges]
            except Exception:
                merges = None
        truncated = (sh.nrows > n_rows) or (sh.ncols > n_cols)
        result.append((name, sheet_to_html(grid, merges, truncated)))
    wb.release_resources()
    if len(names) > keep:
        rest = '、'.join(names[keep: keep + 40])
        more = '' if len(names) - keep <= 40 else f' 等 {len(names) - keep} 个'
        result.append(('__note__', f"<p class='trunc-note'>另有 {len(names) - keep} 个工作表未在此预览：{esc(rest)}{esc(more)}。完整内容请下载原文件查看。</p>"))
    return result


def read_xlsx_sheets(path: str) -> list[tuple[str, str]]:
    keep_merges = os.path.getsize(path) < 4 * 1024 * 1024
    wb = openpyxl.load_workbook(path, read_only=not keep_merges, data_only=True)
    names = wb.sheetnames
    keep = min(len(names), MAX_SHEETS)
    result: list[tuple[str, str]] = []
    for idx in range(keep):
        name = names[idx]
        sh = wb[name]
        n_rows = min(sh.max_row or 0, MAX_ROWS)
        n_cols = min(sh.max_column or 0, MAX_COLS)
        grid = []
        for row in sh.iter_rows(min_row=1, max_row=n_rows, max_col=n_cols, values_only=True):
            grid.append([cell_to_text(v) for v in row])
        grid = trim_grid(grid)
        merges = None
        if keep_merges:
            try:
                merges = [(m.min_row - 1, m.max_row - 1, m.min_col - 1, m.max_col - 1)
                          for m in sh.merged_cells.ranges]
            except Exception:
                merges = None
        truncated = ((sh.max_row or 0) > n_rows) or ((sh.max_column or 0) > n_cols)
        result.append((name, sheet_to_html(grid, merges, truncated)))
    wb.close()
    if len(names) > keep:
        rest = '、'.join(names[keep: keep + 40])
        more = '' if len(names) - keep <= 40 else f' 等 {len(names) - keep} 个'
        result.append(('__note__', f"<p class='trunc-note'>另有 {len(names) - keep} 个工作表未在此预览：{esc(rest)}{esc(more)}。完整内容请下载原文件查看。</p>"))
    return result


# --------------------------------------------------------------------------- #
# Word (.docx 高保真 / .doc 分片表解析) → HTML
# --------------------------------------------------------------------------- #

def iter_docx_blocks(doc: Document):
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag.endswith('}p'):
            yield DocxParagraph(child, doc)
        elif child.tag.endswith('}tbl'):
            yield DocxTable(child, doc)


def docx_to_html(path: str) -> str:
    doc = Document(path)
    out: list[str] = ["<div class='prose'>"]
    for block in iter_docx_blocks(doc):
        if isinstance(block, DocxParagraph):
            text = clean_text(block.text)
            if not text:
                continue
            style = (block.style.name or '') if block.style is not None else ''
            m = re.search(r'(\d+)', style)
            if style.startswith(('Heading', '标题')) and m:
                lvl = min(int(m.group(1)), 4)
                out.append(f'<h{lvl}>{esc(text)}</h{lvl}>')
            else:
                out.append(f'<p>{esc(text)}</p>')
        else:  # table
            rows = []
            for row in block.rows:
                rows.append([clean_text(c.text) for c in row.cells])
            rows = trim_grid(rows)
            out.append(sheet_to_html(rows, None, False))
    out.append('</div>')
    return ''.join(out)


def read_doc_fib_pieces(path: str) -> str | None:
    """解析 Word 97-2003 分片表，取出正文文本。失败返回 None。"""
    if not olefile.isOleFile(path):
        return None
    ole = olefile.OleFileIO(path)
    try:
        if not ole.exists('WordDocument'):
            return None
        wd = ole.openstream('WordDocument').read()
        if len(wd) < 0x200:
            return None
        flags = int.from_bytes(wd[0x0A:0x0C], 'little')
        table_name = '1Table' if (flags & 0x0200) else '0Table'
        if not ole.exists(table_name):
            table_name = '1Table' if ole.exists('1Table') else '0Table'
        if not ole.exists(table_name):
            return None
        table = ole.openstream(table_name).read()

        fc_clx = int.from_bytes(wd[0x01A2:0x01A6], 'little')
        lcb_clx = int.from_bytes(wd[0x01A6:0x01AA], 'little')
        if lcb_clx == 0 or fc_clx + lcb_clx > len(table):
            return None
        clx = table[fc_clx: fc_clx + lcb_clx]

        # 跳过 Prc 块，定位 Pcdt (0x02)
        pos = 0
        while pos < len(clx):
            if clx[pos] == 0x01:
                if pos + 3 > len(clx):
                    return None
                cb = int.from_bytes(clx[pos + 1:pos + 3], 'little')
                pos += 3 + cb
            elif clx[pos] == 0x02:
                break
            else:
                return None
        if pos >= len(clx) or clx[pos] != 0x02:
            return None
        lcb_pcdt = int.from_bytes(clx[pos + 1:pos + 5], 'little')
        plc = clx[pos + 5: pos + 5 + lcb_pcdt]
        if len(plc) < 4:
            return None

        n_pieces = (len(plc) - 4) // 12
        if n_pieces <= 0:
            return None
        cps = [int.from_bytes(plc[i * 4:i * 4 + 4], 'little') for i in range(n_pieces + 1)]
        pcd_off = (n_pieces + 1) * 4

        chunks: list[str] = []
        for i in range(n_pieces):
            pcd = plc[pcd_off + i * 8: pcd_off + i * 8 + 8]
            if len(pcd) < 8:
                continue
            fc = int.from_bytes(pcd[2:6], 'little')
            compressed = bool(fc & 0x40000000)
            offset = (fc & 0x3FFFFFFF) // 2 if compressed else (fc & 0x3FFFFFFF)
            n_chars = cps[i + 1] - cps[i]
            if n_chars <= 0:
                continue
            if compressed:
                raw = wd[offset: offset + n_chars]
                chunks.append(raw.decode('cp1252', errors='replace'))
            else:
                raw = wd[offset: offset + n_chars * 2]
                chunks.append(raw.decode('utf-16-le', errors='replace'))
        text = ''.join(chunks)
        # Word 用 \r 分段、\x07 标记单元格/行结束
        text = text.replace('\x07', '\t').replace('\x0b', '\n')
        return text
    finally:
        ole.close()


def doc_to_html(path: str) -> str | None:
    text = read_doc_fib_pieces(path)
    if not text:
        return None
    paras = [clean_text(p) for p in re.split(r'[\r\n]+', text)]
    paras = [p for p in paras if p]
    if not paras:
        return None
    out = ["<div class='prose'>"]
    for p in paras:
        out.append(f'<p>{esc(p)}</p>')
    out.append('</div>')
    return ''.join(out)


# --------------------------------------------------------------------------- #
# PowerPoint → HTML
# --------------------------------------------------------------------------- #

# 母版 / 占位符噪声：模板自带的「单击此处编辑母版标题样式」这类文字对阅读毫无价值。
_PPT_NOISE = re.compile(
    r'^(单击此处|点击此处|第[一二三四五六七八九十]级$|\*+$|_{2,}P+PT\d*$|'
    r'Click to edit|Second Level|Third Level|Fourth Level|Fifth Level)'
)


def _ppt_line_ok(line: str) -> bool:
    if not line:
        return False
    if line in ('*', '·'):
        return False
    return not _PPT_NOISE.match(line)


def pptx_to_html(path: str) -> str:
    prs = Presentation(path)
    out: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        blocks: list[str] = []
        title = ''
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines = [clean_text(p.text) for p in shape.text_frame.paragraphs]
                lines = [l for l in lines if _ppt_line_ok(l)]
                if not lines:
                    continue
                is_title = False
                try:
                    is_title = shape == slide.shapes.title
                except Exception:
                    is_title = False
                if is_title and not title:
                    title = ' '.join(lines)
                else:
                    blocks.append('<ul>' + ''.join(f'<li>{esc(l)}</li>' for l in lines) + '</ul>')
            elif getattr(shape, 'has_table', False) and shape.has_table:
                rows = [[clean_text(c.text) for c in row.cells] for row in shape.table.rows]
                blocks.append(sheet_to_html(trim_grid(rows), None, False))
        if not title and not blocks:
            continue
        out.append("<section class='slide'>")
        out.append(f"<div class='slide-no'>第 {i} 页</div>")
        if title:
            out.append(f"<h3>{esc(title)}</h3>")
        out.extend(blocks)
        out.append('</section>')
    if not out:
        return "<p class='empty-sheet'>（未解析到可展示的文本内容，请下载原文件查看）</p>"
    return ''.join(out)


def read_ppt_text(path: str) -> list[str] | None:
    """扫描 PowerPoint 97-2003 Document 流里的文本原子，按出现顺序返回文本块。"""
    if not olefile.isOleFile(path):
        return None
    ole = olefile.OleFileIO(path)
    try:
        stream = None
        for name in ('PowerPoint Document', 'PP40'):
            if ole.exists(name):
                stream = ole.openstream(name).read()
                break
        if stream is None:
            return None
        texts: list[str] = []
        pos = 0
        n = len(stream)
        while pos + 8 <= n:
            ver_inst = int.from_bytes(stream[pos:pos + 2], 'little')
            rec_type = int.from_bytes(stream[pos + 2:pos + 4], 'little')
            rec_len = int.from_bytes(stream[pos + 4:pos + 8], 'little')
            rec_ver = ver_inst & 0x000F
            if rec_len < 0 or pos + 8 + rec_len > n:
                break
            if rec_ver == 0xF:  # container
                pos += 8
                continue
            body = stream[pos + 8: pos + 8 + rec_len]
            if rec_type == 0x0FA0:  # TextCharsAtom (UTF-16LE)
                texts.append(body.decode('utf-16-le', errors='replace'))
            elif rec_type == 0x0FA8:  # TextBytesAtom (ANSI)
                texts.append(body.decode('cp1252', errors='replace'))
            elif rec_type == 0x0FBA:  # CString
                texts.append(body.decode('utf-16-le', errors='replace'))
            pos += 8 + rec_len
        out: list[str] = []
        for t in texts:
            t = clean_text(t.replace('\x0b', '\n').replace('\x0d', '\n'))
            for line in t.split('\n'):
                line = clean_text(line)
                if _ppt_line_ok(line):
                    out.append(line)
        return out or None
    finally:
        ole.close()


def ppt_to_html(path: str) -> str | None:
    lines = read_ppt_text(path)
    if not lines:
        return None
    out = ["<div class='prose'>"]
    for l in lines:
        out.append(f'<p>{esc(l)}</p>')
    out.append('</div>')
    return ''.join(out)


# --------------------------------------------------------------------------- #
# 预览页外壳
# --------------------------------------------------------------------------- #

PAGE_CSS = """
:root{color-scheme:light}
*{box-sizing:border-box}
body{margin:0;padding:0;background:#f5f6f8;color:#1d2129;
 font:14px/1.6 -apple-system,BlinkMacSystemFont,"Segoe UI","PingFang SC","Hiragino Sans GB","Microsoft YaHei",sans-serif;}
.wrap{max-width:1400px;margin:0 auto;padding:20px 24px 48px}
.hd{background:#fff;border:1px solid #e5e6eb;border-radius:12px;padding:16px 20px;margin-bottom:16px}
.hd h1{margin:0 0 6px;font-size:17px;font-weight:600;line-height:1.4}
.hd .meta{font-size:12px;color:#86909c;display:flex;flex-wrap:wrap;gap:12px}
.hd .meta b{font-weight:500;color:#4e5969}
.hint{margin-top:10px;font-size:12px;color:#86909c;background:#f7f8fa;border-radius:8px;padding:8px 10px}
.card{background:#fff;border:1px solid #e5e6eb;border-radius:12px;padding:16px 20px;margin-bottom:16px}
/* 工作表切换：纯 CSS（radio + label），页面因此完全不需要脚本，
   宿主可以用最严格的 iframe sandbox 隔离。 */
.tab-radio{position:absolute;width:1px;height:1px;opacity:0;pointer-events:none}
.tabs{display:flex;flex-wrap:wrap;gap:6px;margin:0 0 12px}
.tabs label{font:inherit;font-size:12px;padding:4px 12px;border-radius:999px;border:1px solid #e5e6eb;
 background:#fff;color:#4e5969;cursor:pointer;user-select:none}
.tabs label:hover{border-color:#c9cdd4;background:#f7f8fa}
.pane{display:none}
.pane.on{display:block}
.sheet-wrap{overflow:auto;border:1px solid #e5e6eb;border-radius:8px;max-height:70vh}
table.sheet-table{border-collapse:collapse;font-size:12.5px;min-width:100%;background:#fff}
table.sheet-table td{border:1px solid #e5e6eb;padding:5px 8px;vertical-align:top;
 white-space:pre-wrap;word-break:break-word;min-width:56px;max-width:420px}
table.sheet-table tr:first-child td{background:#f2f3f5;font-weight:600;position:sticky;top:0;z-index:1}
.empty-sheet,.trunc-note{font-size:12px;color:#86909c;margin:8px 0}
.trunc-note{background:#fff7e8;color:#a8661a;border-radius:8px;padding:8px 10px}
h2.sh{font-size:14px;font-weight:600;margin:18px 0 8px;padding-left:8px;border-left:3px solid #165dff}
.prose{font-size:14px;line-height:1.85}
.prose p{margin:0 0 10px}
.prose h1,.prose h2,.prose h3,.prose h4{margin:18px 0 8px;font-weight:600;line-height:1.45}
.prose h1{font-size:18px}.prose h2{font-size:16px}.prose h3{font-size:15px}.prose h4{font-size:14px}
.prose ul{margin:6px 0 12px;padding-left:22px}
.prose li{margin:2px 0}
.slide{border:1px solid #e5e6eb;border-radius:10px;padding:14px 18px;margin-bottom:12px;background:#fff}
.slide-no{font-size:11px;color:#86909c;margin-bottom:4px;letter-spacing:.5px}
.slide h3{margin:0 0 8px;font-size:15px;font-weight:600}
.slide ul{margin:6px 0 0;padding-left:20px}
"""

def build_preview_page(title: str, meta_html: str, hint: str, body: str, extra_css: str = '') -> str:
    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<meta name="referrer" content="no-referrer">
<title>{esc(title)}</title>
<style>{PAGE_CSS}{extra_css}</style>
</head>
<body>
<div class="wrap">
  <div class="hd">
    <h1>{esc(title)}</h1>
    <div class="meta">{meta_html}</div>
    <div class="hint">{hint}</div>
  </div>
  {body}
</div>
</body>
</html>
"""


def sheets_to_body(sheets: list[tuple[str, str]]) -> tuple[str, str]:
    """多工作表 → 纯 CSS 选项卡。返回 (body_html, extra_css)。"""
    notes = [h for (n, h) in sheets if n == '__note__']
    real = [(n, h) for (n, h) in sheets if n != '__note__']
    if not real:
        return (''.join(notes) or "<p class='empty-sheet'>（无可展示内容）</p>"), ''

    if len(real) == 1:
        return (
            '<div class="card">'
            + f'<h2 class="sh">{esc(real[0][0])}</h2>'
            + f'<div class="sheet-wrap">{real[0][1]}</div>'
            + '</div>' + ''.join(notes)
        ), ''

    radios: list[str] = []
    labels: list[str] = []
    panes: list[str] = []
    sel: list[str] = []
    for i, (name, html) in enumerate(real):
        checked = ' checked' if i == 0 else ''
        radios.append(f'<input type="radio" name="sht" id="t{i}" class="tab-radio"{checked}>')
        labels.append(f'<label for="t{i}">{esc(name)}</label>')
        panes.append(f'<div class="pane" id="p{i}"><div class="sheet-wrap">{html}</div></div>')
        sel.append(f'#t{i}:checked~.tabs label[for="t{i}"]')
    extra_css = (
        ','.join(sel) + '{background:#e8f3ff;border-color:#165dff;color:#165dff}'
        + ','.join(f'#t{i}:checked~#p{i}' for i in range(len(real))) + '{display:block}'
    )
    body = (
        '<div class="card">'
        + ''.join(radios)
        + '<div class="tabs">' + ''.join(labels) + '</div>'
        + ''.join(panes)
        + '</div>' + ''.join(notes)
    )
    return body, extra_css


KIND_LABEL = {
    'excel': 'Excel 表格',
    'word': 'Word 文档',
    'ppt': 'PPT 演示',
}

MIME_HINT = {
    'excel': '表格结构已按工作表还原为网页表格，可在线查阅；如需编辑填入数据，请下载原文件。',
    'word': '文档正文已还原为网页排版，可在线查阅；如需编辑，请下载原文件。',
    'ppt': '演示文稿已按页提取文字内容，可在线速览；完整版式请下载原文件。',
}


def human_size(n: int) -> str:
    if n < 1024:
        return f'{n} B'
    if n < 1024 * 1024:
        return f'{n / 1024:.0f} KB'
    return f'{n / 1024 / 1024:.1f} MB'


# --------------------------------------------------------------------------- #
# 主流程
# --------------------------------------------------------------------------- #

def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument('--source', required=True, help='知识库源目录（解压后的文件夹）')
    ap.add_argument('--out', required=True, help='输出目录，如 ui/public/trade-kb')
    args = ap.parse_args()

    src = os.path.abspath(args.source)
    out = os.path.abspath(args.out)
    if not os.path.isdir(src):
        print(f'[ERROR] 源目录不存在: {src}', file=sys.stderr)
        return 2

    files_dir = os.path.join(out, 'files')
    prev_dir = os.path.join(out, 'preview')
    for d in (out, files_dir, prev_dir):
        os.makedirs(d, exist_ok=True)

    names = sorted(os.listdir(src))
    names = [n for n in names if os.path.isfile(os.path.join(src, n))]

    entries: list[dict] = []
    total_bytes = 0
    used_categories: set[str] = set()

    for name in names:
        path = os.path.join(src, name)
        size = os.path.getsize(path)
        total_bytes += size
        ext = os.path.splitext(name)[1].lower().lstrip('.')
        fid = hashlib.md5(name.encode('utf-8')).hexdigest()[:10]

        if ext in ('xls', 'xlsx'):
            kind = 'excel'
        elif ext in ('doc', 'docx'):
            kind = 'word'
        elif ext in ('ppt', 'pptx'):
            kind = 'ppt'
        else:
            kind = 'other'

        cat, summary = FILE_META.get(name, ('reference', ''))
        used_categories.add(cat)

        # 1) 原文件（下载用）
        shutil.copy2(path, os.path.join(files_dir, name))

        # 2) 预览
        preview_rel = None
        preview_note = ''
        body = None
        extra_css = ''
        try:
            if ext == 'xls':
                body, extra_css = sheets_to_body(read_xls_sheets(path))
            elif ext == 'xlsx':
                body, extra_css = sheets_to_body(read_xlsx_sheets(path))
            elif ext == 'docx':
                body = docx_to_html(path)
            elif ext == 'doc':
                body = doc_to_html(path)
                if body is None:
                    preview_note = '旧版 .doc 二进制格式未能完整解析，请下载后查看。'
            elif ext == 'pptx':
                body = pptx_to_html(path)
            elif ext == 'ppt':
                body = ppt_to_html(path)
                if body is None:
                    preview_note = '旧版 .ppt 二进制格式未能完整解析，请下载后查看。'
            else:
                preview_note = '该格式不支持在线预览，请下载后查看。'
        except Exception as e:  # 单文件失败不能拖垮整体
            body = None
            extra_css = ''
            preview_note = f'预览生成失败（{type(e).__name__}），请下载后查看。'
            print(f'[WARN] preview failed: {name} -> {e}', file=sys.stderr)

        if body:
            meta_html = (
                f'<span><b>类型</b> {esc(KIND_LABEL.get(kind, ext.upper()))}</span>'
                f'<span><b>大小</b> {esc(human_size(size))}</span>'
                f'<span><b>文件</b> {esc(name)}</span>'
            )
            hint = MIME_HINT.get(kind, '请下载后查看。')
            if preview_note:
                hint = f'{hint} {preview_note}'
            page = build_preview_page(name, meta_html, hint, body, extra_css)
            prev_file = f'{fid}.html'
            # 必须带 newline=''：否则 Python 在 Windows 上会把 LF 写成 CRLF，
            # 违反本仓 .gitattributes 的 `* text=auto eol=lf`，每次构建都产生无谓 diff。
            with open(os.path.join(prev_dir, prev_file), 'w', encoding='utf-8', newline='') as f:
                f.write(page)
            preview_rel = f'preview/{prev_file}'
        else:
            preview_note = preview_note or '暂不支持在线预览，请下载后查看。'

        entries.append({
            'id': fid,
            'name': name,
            'ext': ext,
            'kind': kind,
            'size': size,
            'sizeText': human_size(size),
            'category': cat,
            'summary': summary,
            'href': 'files/' + _quote(name),
            'preview': preview_rel,
            'previewNote': preview_note,
        })
        print(f'  [{cat:13s}] {name}  {"preview ok" if preview_rel else "download only"}')

    # 分类排序：只保留实际用到的分类
    cats = [c for c in CATEGORIES if c['id'] in used_categories]
    order = {c['id']: c['order'] for c in CATEGORIES}
    entries.sort(key=lambda e: (order.get(e['category'], 99), e['name']))

    manifest = {
        'version': 1,
        'source': '外贸人知识库',
        'generatedAt': _dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'stats': {
            'files': len(entries),
            'bytes': total_bytes,
            'bytesText': human_size(total_bytes),
            'categories': len(cats),
            'previewable': sum(1 for e in entries if e['preview']),
        },
        'categories': cats,
        'files': entries,
    }
    with open(os.path.join(out, 'manifest.json'), 'w', encoding='utf-8', newline='') as f:
        json.dump(manifest, f, ensure_ascii=False, indent=1)

    print(f'\nOK  {len(entries)} 个文件 / {human_size(total_bytes)} / {len(cats)} 个分类')
    print(f'    可在线预览 {manifest["stats"]["previewable"]} 个')
    print(f'    输出 {out}')
    return 0


def _quote(name: str) -> str:
    from urllib.parse import quote
    return quote(name, safe='')


if __name__ == '__main__':
    raise SystemExit(main())
